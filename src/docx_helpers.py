"""Shared python-docx helpers (imported by the report builders via runpy-free exec; see 05_build_reports.py)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pandas as pd

BLUE = RGBColor(0x1f,0x4e,0x8c); INK = RGBColor(0x1a,0x1a,0x1a); GREY = RGBColor(0x55,0x55,0x55)

def new_doc(header_text):
    d = Document(); d.core_properties.author = "Samatar Ahmed"; d.core_properties.title = header_text.split(" | ")[0]; d.core_properties.last_modified_by = "Samatar Ahmed"; s = d.sections[0]; s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.4); s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.18
    for n, sz, col in (("Heading 1", 18, BLUE), ("Heading 2", 13.5, BLUE), ("Heading 3", 11.5, INK)):
        h = d.styles[n]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = col
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); h.paragraph_format.space_before = Pt(14 if n != "Heading 3" else 10); h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    # header / footer with page number
    hp = s.header.paragraphs[0]; hp.text = header_text; hp.runs[0].font.size = Pt(8.5); hp.runs[0].font.color.rgb = GREY
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(); r.font.size = Pt(9)
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)
    return d

def h1(d, t, newpage=True):
    if newpage: d.add_page_break()
    return d.add_heading(t, 1)
def h2(d, t): return d.add_heading(t, 2)
def h3(d, t): return d.add_heading(t, 3)

def para(d, text, bold_lead=None, italic=False, size=None, align=None, color=None, after=None):
    p = d.add_paragraph()
    if bold_lead: r = p.add_run(bold_lead + " "); r.bold = True
    r = p.add_run(text); r.italic = italic
    if size: [setattr(x.font, "size", Pt(size)) for x in p.runs]
    if color: [setattr(x.font, "color", color) for x in p.runs] if False else [setattr(x.font.color, "rgb", color) for x in p.runs]
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def bullets(d, items, style="List Bullet"):
    for it in items:
        p = d.add_paragraph(style=style); p.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple): r = p.add_run(it[0] + " "); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)

def callout(d, title, text, fill="EAF1FB"):
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; c = t.rows[0].cells[0]; shade(c, fill)
    c.paragraphs[0].add_run(title).bold = True; p = c.add_paragraph(text); p.paragraph_format.space_after = Pt(2)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

def table(d, df, caption=None, widths=None, fs=9, align_right_from=1):
    if caption:
        p = d.add_paragraph(); r = p.add_run(caption); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = GREY; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(3)
    t = d.add_table(rows=1, cols=len(df.columns)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(df.columns):
        cell = t.rows[0].cells[i]; cell.text = ""; r = cell.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(fs); r.font.color.rgb = RGBColor(255,255,255); shade(cell, "1F4E8C")
    for ri, row in enumerate(df.itertuples(index=False)):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; p = cells[i].paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(fs)
            if i >= align_right_from: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(0)
            if ri % 2: shade(cells[i], "F3F5F8")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

_fig = [0]
def figure(d, path, caption, width=15.2):
    _fig[0] += 1
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path, width=Cm(width)); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
    c = d.add_paragraph(); r = c.add_run(f"Figure {_fig[0]}. {caption}"); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY; c.alignment = WD_ALIGN_PARAGRAPH.CENTER

def code(d, text, fs=8):
    for line in text.strip("\n").split("\n"):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " "); r.font.name = "Consolas"; r.font.size = Pt(fs); r._r.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    d.add_paragraph().paragraph_format.space_after = Pt(2)
